# app/services/document_parser.py
import os
import re
import json
import logging
from typing import Dict, Any, List, Optional
from pathlib import Path
from datetime import datetime

from docx import Document
import pdfplumber
import PyPDF2

# 新增导入LangChain分割器
from app.utils.langchain_splitter import get_splitter_for_ct_report

logger = logging.getLogger(__name__)


class DocumentParser:
    """文档解析服务 - 从Word/PDF中提取RAG知识库数据"""

    def __init__(self, enable_chunking: bool = True):
        """
        初始化文档解析器

        Args:
            enable_chunking: 是否启用文本分割（默认启用）
        """
        self.enable_chunking = enable_chunking
        if enable_chunking:
            self.splitter = get_splitter_for_ct_report()
            logger.info("✓ 文档分割器已启用")

        self.section_markers = {
            "case_id": "病例编号:",
            "task": "任务类型:",
            "summary": "影像摘要:",
            "positive_pixels": "阳性像素数:",
            "total_pixels": "总像素数:",
            "ratio": "阳性占比:",
            "slice_indices": "涉及切片:",
            "nodule_type": "结节类型:",
            "size_mm": "结节大小:",
            "location": "结节位置:",
            "edge": "边缘特征:",
            "ai_summary": "AI诊断结论:",
            "risk_level": "风险等级:",
            "suggestions": "处理建议:",
            "follow_up": "随访建议:",
            "doctor_diagnosis": "最终诊断:",
            "treatment": "治疗方案:",
            "pathology": "病理结果:",
            "confirmed_by": "确诊医生:",
            "follow_up_months": "随访月数:",
            "treatment_effectiveness": "治疗效果:",
            "recurrence": "是否复发:",
            "quality_score": "质量评分:",
            "hospital": "医院:"
        }

    def parse_document(self, file_path: str) -> List[Dict[str, Any]]:
        """
        解析文档文件，返回一个或多个案例

        Args:
            file_path: 文档路径 (.docx 或 .pdf)

        Returns:
            案例列表（每个案例可能包含多个文本块）
        """
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")

        # 1. 提取文本
        ext = file_path.suffix.lower()
        if ext == '.docx':
            text = self._parse_docx(file_path)
        elif ext == '.pdf':
            text = self._parse_pdf(file_path)
        else:
            raise ValueError(f"不支持的文件类型: {ext}")

        # 2. 解析文本提取数据
        return self._parse_text_to_cases(text, file_path.stem)

    def _parse_docx(self, file_path: Path) -> str:
        """解析Word文档"""
        try:
            doc = Document(file_path)
            text_parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

            # 读取表格
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells])
                    if row_text.strip():
                        text_parts.append(row_text)

            return '\n'.join(text_parts)
        except Exception as e:
            logger.error(f"Word文档解析失败: {str(e)}")
            raise

    def _parse_pdf(self, file_path: Path) -> str:
        """解析PDF文档"""
        try:
            text_parts = []

            # 优先使用pdfplumber
            try:
                with pdfplumber.open(file_path) as pdf:
                    for page in pdf.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text_parts.append(page_text)
            except:
                # 降级到PyPDF2
                with open(file_path, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    for page in reader.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text_parts.append(page_text)

            return '\n'.join(text_parts)
        except Exception as e:
            logger.error(f"PDF文档解析失败: {str(e)}")
            raise

    def _parse_text_to_cases(self, text: str, filename: str) -> List[Dict[str, Any]]:
        """
        从文本解析案例数据

        如果文档包含多个案例（用"---"分隔），会分别解析
        """
        # 检查是否包含多个案例
        case_separator = "---"
        if case_separator in text:
            case_texts = text.split(case_separator)
            logger.info(f"检测到 {len(case_texts)} 个案例")
        else:
            case_texts = [text]

        cases = []
        for idx, case_text in enumerate(case_texts):
            if not case_text.strip():
                continue

            # 提取字段
            data = {}
            for key, marker in self.section_markers.items():
                data[key] = self._extract_value(case_text, marker)

            # 生成案例ID
            case_id = data.get("case_id") or f"CASE_{filename}_{idx + 1}"

            # 构建标准案例数据
            case_data = self._build_case_data(case_id, data)

            # 如果需要分割，分割summary并创建多个块
            if self.enable_chunking and case_data.get('summary'):
                cases.extend(self._split_case_into_chunks(case_data))
            else:
                cases.append(case_data)

        return cases

    def _build_case_data(self, case_id: str, data: Dict[str, Any]) -> Dict[str, Any]:
        """构建案例数据"""
        summary = data.get("summary", "")

        return {
            "case_id": case_id,
            "model_version": "CT-V3.2",
            "hospital": data.get("hospital", "未知医院"),
            "is_active": True,
            "summary": summary,  # 完整摘要
            "model_features": {
                "task": data.get("task", "肺结节检测"),
                "summary": summary,
                "positive_pixels": self._to_int(data.get("positive_pixels", 0)),
                "total_pixels": self._to_int(data.get("total_pixels", 100000)),
                "ratio": self._to_float(data.get("ratio", 0)),
                "slice_indices": self._to_list(data.get("slice_indices", "")),
                "nodule_type": data.get("nodule_type", ""),
                "size_mm": self._to_float(data.get("size_mm", 0)),
                "location": data.get("location", ""),
                "edge": data.get("edge", "")
            },
            "ai_analysis": {
                "summary": data.get("ai_summary", summary),
                "riskLevel": data.get("risk_level", "MEDIUM"),
                "suggestions": self._to_list_items(data.get("suggestions", "")),
                "followUpAdvice": data.get("follow_up", "建议3个月后复查")
            },
            "doctor_confirmed": {
                "diagnosis": data.get("doctor_diagnosis", ""),
                "treatment": data.get("treatment", ""),
                "pathology_result": data.get("pathology", ""),
                "confirmed_by": data.get("confirmed_by", ""),
                "confirmed_at": datetime.now().isoformat()
            },
            "validation": {
                "is_pathology_confirmed": bool(data.get("pathology", "")),
                "follow_up_months": self._to_int(data.get("follow_up_months", 0)),
                "treatment_effectiveness": data.get("treatment_effectiveness", ""),
                "recurrence": self._to_bool(data.get("recurrence", "false")),
                "quality_score": self._to_float(data.get("quality_score", 0.7))
            },
            "search_text": summary  # 直接用摘要作为检索文本
        }

    def _split_case_into_chunks(self, case_data: Dict[str, Any]) -> List[Dict[str, Any]]:
        """
        将案例分割成多个文本块

        每个块包含部分摘要内容，但保留完整的元数据
        """
        summary = case_data.get('summary', '')
        if not summary or len(summary) <= self.splitter.chunk_size:
            # 如果摘要很短，不需要分割
            return [case_data]

        # 分割摘要
        chunks = self.splitter.split_text(summary)
        logger.info(f"案例 {case_data.get('case_id')} 分割为 {len(chunks)} 个块")

        # 为每个块创建独立的案例
        result = []
        for idx, chunk in enumerate(chunks):
            chunk_case = case_data.copy()
            chunk_case['case_id'] = f"{case_data.get('case_id')}_chunk_{idx + 1}"
            chunk_case['summary'] = chunk
            chunk_case['model_features'] = case_data['model_features'].copy()
            chunk_case['model_features']['summary'] = chunk
            chunk_case['ai_analysis'] = case_data['ai_analysis'].copy()
            chunk_case['ai_analysis']['summary'] = chunk
            chunk_case['search_text'] = chunk
            chunk_case['chunk_index'] = idx
            chunk_case['chunk_total'] = len(chunks)
            chunk_case['parent_case_id'] = case_data.get('case_id')

            result.append(chunk_case)

        return result

    def _extract_value(self, text: str, marker: str) -> str:
        """从文本中提取标记后面的值"""
        pattern = f"{marker}(.*?)(?=\\n\\n|\\Z)"
        match = re.search(pattern, text, re.DOTALL | re.IGNORECASE)
        if match:
            return match.group(1).strip()
        return ""

    def _to_int(self, value: Any) -> int:
        """转换为整数"""
        if isinstance(value, int):
            return value
        if isinstance(value, str):
            nums = re.findall(r'\d+', value)
            return int(nums[0]) if nums else 0
        return 0

    def _to_float(self, value: Any) -> float:
        """转换为浮点数"""
        if isinstance(value, (int, float)):
            return float(value)
        if isinstance(value, str):
            nums = re.findall(r'[\d.]+', value)
            return float(nums[0]) if nums else 0.0
        return 0.0

    def _to_bool(self, value: Any) -> bool:
        """转换为布尔值"""
        if isinstance(value, bool):
            return value
        if isinstance(value, str):
            return value.lower() in ['true', '是', 'yes', '1', 't']
        return False

    def _to_list(self, value: str) -> List[int]:
        """转换为列表"""
        if not value:
            return []
        items = [v.strip() for v in value.split(',') if v.strip()]
        return [int(v) for v in items if v.isdigit()]

    def _to_list_items(self, text: str) -> List[str]:
        """提取列表项"""
        if not text:
            return ["建议由主治医生综合评估"]
        items = [v.strip() for v in text.split('\n') if v.strip()]
        items = [re.sub(r'^[-•*]\s*', '', item).strip() for item in items]
        items = [re.sub(r'^\d+[.、]\s*', '', item).strip() for item in items]
        return items if items else ["建议由主治医生综合评估"]


class BatchDocumentParser:
    """批量文档解析器"""

    def __init__(self, enable_chunking: bool = True):
        self.parser = DocumentParser(enable_chunking=enable_chunking)

    def parse_directory(self, directory_path: str, recursive: bool = True) -> List[Dict[str, Any]]:
        """
        批量解析目录下的所有文档

        Args:
            directory_path: 目录路径
            recursive: 是否递归子目录

        Returns:
            解析后的案例列表
        """
        directory = Path(directory_path)
        if not directory.exists():
            raise FileNotFoundError(f"目录不存在: {directory_path}")

        supported_extensions = ['.docx', '.pdf']
        all_cases = []

        pattern = "**/*" if recursive else "*"
        for file_path in directory.glob(pattern):
            if file_path.is_file() and file_path.suffix.lower() in supported_extensions:
                try:
                    cases = self.parser.parse_document(str(file_path))
                    all_cases.extend(cases)
                    logger.info(f"✓ 解析文档: {file_path.name} -> {len(cases)} 个案例")
                except Exception as e:
                    logger.error(f"✗ 解析失败: {file_path.name} - {str(e)}")

        logger.info(f"批量解析完成: 共 {len(all_cases)} 个案例")
        return all_cases